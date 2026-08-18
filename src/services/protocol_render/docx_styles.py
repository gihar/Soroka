"""Деловой облик Word-протокола: геометрия страницы, стили Word, колонтитул.

Разделение с ``docx_renderer`` намеренное: рендерер разбирает канонический
Markdown и раскладывает его по стилям, этот модуль отвечает за то, как эти стили
выглядят. Правка облика не трогает разбор, правка разбора не трогает облик.

Облик задан ПЕРЕОПРЕДЕЛЕНИЕМ встроенных стилей Word (Normal, Heading 1-4,
List Number, List Bullet), а не прямым форматированием runs. Это принципиально:
протокол в Word открывают, чтобы дописать пункт и переслать дальше. Абзац,
дописанный пользователем, обязан выглядеть как остальные, а «Изменить стиль»
обязан менять документ целиком — прямое форматирование дало бы картинку, которую
нельзя править (ADR-0009).
"""

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from src.utils import document_palette as palette

# Arial — единственная гротеска, которая есть в Word и на Windows, и на macOS, и
# предсказуемо подменяется в LibreOffice (Liberation Sans). Calibri читается как
# «документ никто не оформлял» (это шрифт по умолчанию), Times New Roman тянет
# канцелярскую интонацию, которой продукт избегает (анти-референс PRODUCT.md).
FONT = "Arial"

# Стиль шапки: дата и участники под заголовком. Имя по-русски — пользователь
# видит его в панели стилей Word.
META_STYLE = "Шапка протокола"

_TITLE_PT = 18
_H2_PT = 12
_H3_PT = 11
_H4_PT = 10.5
_BODY_PT = 10.5
_META_PT = 9.5
_FOOTER_PT = 8

# Порядок детей внутри w:pPr и w:rPr задан схемой OOXML; python-docx свои
# последовательности удаляет после определения класса, поэтому нужные хвосты
# перечислены здесь. Вставка не по порядку даёт документ, который Word считает
# повреждённым.
_PBDR_SUCCESSORS = ("w:shd", "w:tabs", "w:spacing", "w:ind", "w:jc", "w:outlineLvl")
_CONTEXTUAL_SPACING_SUCCESSORS = ("w:mirrorIndents", "w:suppressOverlap", "w:jc",
                                  "w:outlineLvl", "w:rPr", "w:sectPr")
_LANG_SUCCESSORS = ("w:eastAsianLayout", "w:specVanish", "w:oMath")

# Тема (asciiTheme/themeColor) в стилях Word по умолчанию перебивает явные
# значения — снимаем её, иначе заголовки останутся синим Calibri Light.
_THEME_FONT_ATTRS = ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme")
_THEME_COLOR_ATTRS = ("themeColor", "themeShade", "themeTint")


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#").upper())


def _drop_theme_font(rpr) -> None:
    """Снять тематические шрифты и прописать явные для кириллицы и CS."""
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        return
    for attr in _THEME_FONT_ATTRS:
        rfonts.attrib.pop(qn(f"w:{attr}"), None)
    rfonts.set(qn("w:cs"), FONT)
    rfonts.set(qn("w:eastAsia"), FONT)


def _drop_theme_color(rpr) -> None:
    color = rpr.find(qn("w:color"))
    if color is None:
        return
    for attr in _THEME_COLOR_ATTRS:
        color.attrib.pop(qn(f"w:{attr}"), None)


def _apply_font(style, *, size_pt: float, color: str, bold: bool | None = None) -> None:
    """Назначить стилю шрифт, кегль и цвет, вытеснив тему документа."""
    font = style.font
    font.name = FONT
    font.size = Pt(size_pt)
    font.color.rgb = _rgb(color)
    if bold is not None:
        font.bold = bold
    rpr = style.element.get_or_add_rPr()
    _drop_theme_font(rpr)
    _drop_theme_color(rpr)
    # Кегль сложных письменностей (szCs) во встроенных стилях свой; оставленный
    # без внимания, он рассинхронизирует заголовок на смешанном тексте.
    size_cs = rpr.find(qn("w:szCs"))
    if size_cs is not None:
        size_cs.set(qn("w:val"), str(int(size_pt * 2)))


def _set_russian_language(style) -> None:
    """Пометить текст русским: без этого Word проверяет протокол как английский."""
    rpr = style.element.get_or_add_rPr()
    if rpr.find(qn("w:lang")) is not None:
        return
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "ru-RU")
    rpr.insert_element_before(lang, *_LANG_SUCCESSORS)


def _set_top_rule(style, color: str) -> None:
    """Тонкая линейка над абзацем — разделитель секций (паритет с PDF)."""
    ppr = style.element.get_or_add_pPr()
    if ppr.find(qn("w:pBdr")) is not None:
        return
    borders = OxmlElement("w:pBdr")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), "4")  # восьмые доли пункта: 4 = 0,5 pt
    top.set(qn("w:space"), "8")  # отбивка линейки от текста, пункты
    top.set(qn("w:color"), color.lstrip("#"))
    borders.append(top)
    ppr.insert_element_before(borders, *_PBDR_SUCCESSORS)


def _add_contextual_spacing(style) -> None:
    """Не разводить интервалом абзацы одного стиля — сомкнуть их в блок."""
    ppr = style.element.get_or_add_pPr()
    if ppr.find(qn("w:contextualSpacing")) is not None:
        return
    ppr.insert_element_before(
        OxmlElement("w:contextualSpacing"), *_CONTEXTUAL_SPACING_SUCCESSORS
    )


def _drop_contextual_spacing(style) -> None:
    """Снять «не добавлять интервал между абзацами одного стиля».

    Пункты протокола переносятся на две-три строки; без интервала между ними
    список слипается в кирпич, и глаз перестаёт находить границы пунктов.
    """
    ppr = style.element.get_or_add_pPr()
    contextual = ppr.find(qn("w:contextualSpacing"))
    if contextual is not None:
        ppr.remove(contextual)


def _style_page(section) -> None:
    """A4 и поля 2 см — та же геометрия, что у PDF-канала.

    По умолчанию python-docx отдаёт US Letter: документ, свёрстанный под лист,
    которого нет ни в одной российской переговорке.
    """
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)


def _field_char(kind: str):
    element = OxmlElement("w:fldChar")
    element.set(qn("w:fldCharType"), kind)
    return element


def _add_page_number(section) -> None:
    """«стр. N» в нижнем колонтитуле — поле Word, а не проставленное число.

    Поле собрано полной пятёркой (begin → instrText → separate → результат →
    end): без ветки результата номер берёт оформление не из своего run и
    выпадает из колонтитула другим кеглем.
    """
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    label = paragraph.add_run("стр. ")

    begin = paragraph.add_run()
    begin._r.append(_field_char("begin"))

    instruction = paragraph.add_run()
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    instruction._r.append(instr_text)

    separate = paragraph.add_run()
    separate._r.append(_field_char("separate"))

    result = paragraph.add_run("1")  # значение до пересчёта поля в Word

    end = paragraph.add_run()
    end._r.append(_field_char("end"))

    for run in (label, begin, instruction, separate, result, end):
        run.font.name = FONT
        run.font.size = Pt(_FOOTER_PT)
        run.font.color.rgb = _rgb(palette.COLOR_MUTED)


def _style_body(styles) -> None:
    normal = styles["Normal"]
    _apply_font(normal, size_pt=_BODY_PT, color=palette.COLOR_BODY)
    _set_russian_language(normal)
    fmt = normal.paragraph_format
    fmt.line_spacing = 1.15
    fmt.space_after = Pt(4)
    fmt.widow_control = True


def _style_headings(styles) -> None:
    title = styles["Heading 1"]
    _apply_font(title, size_pt=_TITLE_PT, color=palette.COLOR_TITLE, bold=True)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(6)

    section = styles["Heading 2"]
    _apply_font(section, size_pt=_H2_PT, color=palette.COLOR_HEADING, bold=True)
    section.paragraph_format.space_before = Pt(16)
    section.paragraph_format.space_after = Pt(6)
    section.paragraph_format.keep_with_next = True
    # Линейка сверху закрывает предыдущую секцию и шапку заодно; привязанная к
    # заголовку, она никогда не повиснет одна внизу страницы.
    _set_top_rule(section, palette.COLOR_RULE)

    for name, size in (("Heading 3", _H3_PT), ("Heading 4", _H4_PT)):
        subsection = styles[name]
        _apply_font(subsection, size_pt=size, color=palette.COLOR_MUTED, bold=True)
        subsection.paragraph_format.space_before = Pt(10)
        subsection.paragraph_format.space_after = Pt(2)
        subsection.paragraph_format.keep_with_next = True


def _style_lists(styles) -> None:
    for name in ("List Number", "List Bullet"):
        style = styles[name]
        _apply_font(style, size_pt=_BODY_PT, color=palette.COLOR_BODY)
        _drop_contextual_spacing(style)
        fmt = style.paragraph_format
        fmt.left_indent = Cm(0.75)
        fmt.first_line_indent = Cm(-0.75)  # висячий отступ: номер прижат к полю
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(3)
        fmt.line_spacing = 1.15
        fmt.widow_control = True


def _add_meta_style(document) -> None:
    """Стиль шапки: мельче и тише тела — дата и участники не спорят с решениями."""
    style = document.styles.add_style(META_STYLE, WD_STYLE_TYPE.PARAGRAPH)
    # Имя стиля русское (его видит пользователь в панели стилей), а идентификатор
    # латиницей: python-docx вывел бы id из имени, а кириллический styleId —
    # лишний риск на чужой сборке Word.
    style.style_id = "ProtocolHeader"
    style.base_style = document.styles["Normal"]
    _apply_font(style, size_pt=_META_PT, color=palette.COLOR_MUTED)
    style.paragraph_format.line_spacing = 1.1
    # Отбивка стоит ПОСЛЕ блока, а не после каждой строки: строки шапки смыкаются
    # (дата, лектор, столбик участников — один блок), а от того, что идёт следом,
    # блок отделён. Иначе шапка липнет к тексту — как в «Протоколе ОД», где
    # вместо секции сразу идёт «Поручений в записи не зафиксировано».
    _add_contextual_spacing(style)
    style.paragraph_format.space_after = Pt(8)


def _clear_generator_metadata(document) -> None:
    """Снять подпись библиотеки: протокол — документ владельца встречи.

    python-docx по умолчанию записывает себя в автора и последнего
    редактировавшего; в свойствах пересланного файла это выглядит как служебный
    мусор (анти-референс PRODUCT.md).
    """
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""


def build_protocol_document() -> Document:
    """Пустой документ протокола с деловым стилем — основа для рендерера."""
    document = Document()
    _style_page(document.sections[0])
    _add_page_number(document.sections[0])
    _style_body(document.styles)
    _style_headings(document.styles)
    _style_lists(document.styles)
    _add_meta_style(document)
    _clear_generator_metadata(document)
    return document
