"""Word-протокол в строгом деловом стиле (ADR-0009).

До этой правки Word-канал отдавал документ ровно в том виде, в каком его собирает
python-docx: US Letter вместо A4, синие заголовки Calibri Light темы Word,
эмодзи-метки секций цветными картинками, автор файла «python-docx» и ни одного
номера страницы. Такой файл пересылают «наверх» — и он выглядит выгрузкой из
чата, а не протоколом.

Тесты читают сгенерированный .docx обратно и проверяют наблюдаемое: стили,
геометрию страницы, свойства файла. Проверяется контракт («заголовки не темы
Word», «страница A4»), а не конкретные пункты кегля — облик можно донастраивать,
не переписывая тесты.
"""

import io

import pytest
from docx import Document
from docx.oxml.ns import qn

from src.services.protocol_render.docx_renderer import convert_protocol_to_docx
from src.services.protocol_render.docx_styles import FONT, META_STYLE
from src.utils import document_palette as palette

_PROTOCOL = """# Планёрка по запуску

**Дата:** 18 августа 2026 · 14:30
**👥 Участники:** Ли Ясмина, Иван Петров

## ✅ Решения
1. Бету запускаем 12 ноября.

## ⚠️ Блокеры и риски
1. **Риск**: нагрузочный тест не проведён.
"""


@pytest.fixture(scope="module")
def document():
    return Document(io.BytesIO(convert_protocol_to_docx(_PROTOCOL)))


def _style(document, name):
    return document.styles[name]


def _rpr(style):
    return style.element.get_or_add_rPr()


def _hex(color: str) -> str:
    return color.lstrip("#").upper()


# ---------------------------------------------------------------------------
# 1. Эмодзи: документные каналы идут без них
# ---------------------------------------------------------------------------


def test_section_headings_have_no_emoji(document):
    headings = [p.text for p in document.paragraphs if p.style.name == "Heading 2"]
    assert headings == ["Решения", "Блокеры и риски"]


def test_header_label_keeps_no_gap_where_emoji_was(document):
    участники = next(p for p in document.paragraphs if "Участники" in p.text)
    assert участники.text == "Участники: Ли Ясмина, Иван Петров"


# ---------------------------------------------------------------------------
# 2. Геометрия страницы: A4 и поля как у PDF-канала
# ---------------------------------------------------------------------------


def test_page_is_a4_with_two_centimetre_margins(document):
    # Word хранит размеры в твипах, поэтому сверяем миллиметры, а не EMU.
    section = document.sections[0]
    assert (round(section.page_width.mm), round(section.page_height.mm)) == (210, 297)
    assert round(section.left_margin.mm) == round(section.right_margin.mm) == 20
    assert round(section.top_margin.mm) == round(section.bottom_margin.mm) == 20


def test_footer_carries_a_page_field(document):
    footer_xml = document.sections[0].footer.paragraphs[0]._p.xml
    assert "стр." in footer_xml
    # Именно ПОЛЕ, а не проставленная единица: на второй странице должно быть «2».
    assert "PAGE" in footer_xml
    assert 'w:fldCharType="begin"' in footer_xml
    assert 'w:fldCharType="end"' in footer_xml


# ---------------------------------------------------------------------------
# 3. Типографика: своя, а не тема Word по умолчанию
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "style_name",
    ["Normal", "Heading 1", "Heading 2", "Heading 3", "List Number", "List Bullet"],
)
def test_styles_do_not_fall_back_to_word_theme(document, style_name):
    """Тема перебивает явный шрифт и цвет — заголовки остались бы синим Calibri."""
    rpr = _rpr(_style(document, style_name))
    fonts = rpr.find(qn("w:rFonts"))
    assert fonts is not None, style_name
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        assert fonts.get(qn(f"w:{attr}")) is None, (style_name, attr)
    color = rpr.find(qn("w:color"))
    if color is not None:
        assert color.get(qn("w:themeColor")) is None, style_name


@pytest.mark.parametrize(
    "style_name",
    ["Normal", "Heading 1", "Heading 2", "List Number", "List Bullet", META_STYLE],
)
def test_every_protocol_style_uses_the_document_font(document, style_name):
    assert _style(document, style_name).font.name == FONT


@pytest.mark.parametrize(
    "style_name, expected",
    [
        ("Normal", palette.COLOR_BODY),
        ("Heading 1", palette.COLOR_TITLE),
        ("Heading 2", palette.COLOR_HEADING),
        (META_STYLE, palette.COLOR_MUTED),
    ],
)
def test_styles_take_colour_from_the_shared_document_palette(
    document, style_name, expected
):
    """Палитра общая с PDF: один документ в двух форматах, а не два вида."""
    assert str(_style(document, style_name).font.color.rgb) == _hex(expected)


def test_title_outranks_section_headings_by_size(document):
    title = _style(document, "Heading 1").font.size
    section = _style(document, "Heading 2").font.size
    body = _style(document, "Normal").font.size
    assert title > section > body


def test_section_heading_carries_a_rule_and_stays_with_its_content(document):
    heading = _style(document, "Heading 2")
    ppr = heading.element.get_or_add_pPr()
    borders = ppr.find(qn("w:pBdr"))
    assert borders is not None and borders.find(qn("w:top")) is not None
    assert borders.find(qn("w:top")).get(qn("w:color")) == _hex(palette.COLOR_RULE)
    # Линейка привязана к заголовку — она не повиснет одна внизу страницы,
    # а заголовок не оторвётся от своего списка.
    assert heading.paragraph_format.keep_with_next is True


def test_document_is_typed_as_russian(document):
    lang = _rpr(_style(document, "Normal")).find(qn("w:lang"))
    assert lang is not None and lang.get(qn("w:val")) == "ru-RU"


def test_list_items_are_spaced_apart(document):
    """Пункты переносятся на две строки — без интервала список слипается."""
    for name in ("List Number", "List Bullet"):
        style = _style(document, name)
        assert style.paragraph_format.space_after.pt > 0, name
        contextual = style.element.get_or_add_pPr().find(qn("w:contextualSpacing"))
        assert contextual is None, name


# ---------------------------------------------------------------------------
# 4. Шапка документа отделена от содержания
# ---------------------------------------------------------------------------


def test_header_lines_use_the_meta_style(document):
    styles = {p.text: p.style.name for p in document.paragraphs if p.text}
    assert styles["Дата: 18 августа 2026 · 14:30"] == META_STYLE
    assert styles["Участники: Ли Ясмина, Иван Петров"] == META_STYLE


def test_meta_style_is_quieter_than_body(document):
    assert _style(document, META_STYLE).font.size < _style(document, "Normal").font.size


def test_meta_block_is_tight_inside_and_spaced_after(document):
    """Строки шапки — один блок; отбивка стоит после блока, а не после каждой.

    Иначе шапка липнет к тексту там, где следом идёт не секция, а абзац
    («Протокол ОД» с пустыми поручениями).
    """
    style = _style(document, META_STYLE)
    contextual = style.element.get_or_add_pPr().find(qn("w:contextualSpacing"))
    assert contextual is not None
    assert style.paragraph_format.space_after.pt > 0


def test_content_paragraphs_are_not_meta_styled():
    doc = Document(io.BytesIO(convert_protocol_to_docx(
        "# Титул\n\n**Дата:** сегодня\n\n## Обсуждение\nОбычный абзац.\n"
    )))
    styles = {p.text: p.style.name for p in doc.paragraphs if p.text}
    assert styles["Обычный абзац."] == "Normal"


def test_header_block_closes_on_the_first_blank_line():
    """Нестандартный шаблон без секций не уходит в шапку целиком."""
    doc = Document(io.BytesIO(convert_protocol_to_docx(
        "# Титул\n\n**Дата:** сегодня\n\nПервый абзац заметки.\nВторой абзац.\n"
    )))
    styles = {p.text: p.style.name for p in doc.paragraphs if p.text}
    assert styles["Дата: сегодня"] == META_STYLE
    assert styles["Первый абзац заметки."] == "Normal"
    assert styles["Второй абзац."] == "Normal"


def test_protocol_without_title_has_no_meta_block():
    doc = Document(io.BytesIO(convert_protocol_to_docx("Просто текст без титула.\n")))
    assert all(p.style.name != META_STYLE for p in doc.paragraphs)


# ---------------------------------------------------------------------------
# 5. Свойства файла: документ владельца встречи, а не библиотеки
# ---------------------------------------------------------------------------


def test_document_title_comes_from_the_protocol(document):
    assert document.core_properties.title == "Планёрка по запуску"


def test_generator_signature_is_cleared(document):
    props = document.core_properties
    assert props.author == ""
    assert props.last_modified_by == ""


def test_title_survives_a_protocol_with_several_h1():
    doc = Document(io.BytesIO(convert_protocol_to_docx("# Первый\n\n# Второй\n")))
    assert doc.core_properties.title == "Первый"
